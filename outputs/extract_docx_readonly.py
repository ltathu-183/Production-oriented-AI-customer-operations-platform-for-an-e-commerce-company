from pathlib import Path
import sys

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(parent):
    element = parent.element.body
    for child in element.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def main(path_text: str) -> None:
    path = Path(path_text)
    doc = Document(path)
    print(f"# FILE: {path.name}")
    print(f"# PARAGRAPHS: {len(doc.paragraphs)} TABLES: {len(doc.tables)} SECTIONS: {len(doc.sections)}")
    for index, block in enumerate(iter_blocks(doc), 1):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                print(f"\n[P{index} style={block.style.name!r}] {text}")
        else:
            print(f"\n[T{index} rows={len(block.rows)} cols={len(block.columns)}]")
            for row_index, row in enumerate(block.rows, 1):
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                print(f"  R{row_index}: " + " | ".join(cells))
    for section_index, section in enumerate(doc.sections, 1):
        for label, part in (("HEADER", section.header), ("FOOTER", section.footer)):
            lines = [p.text.strip() for p in part.paragraphs if p.text.strip()]
            if lines:
                print(f"\n[{label} section={section_index}] " + " || ".join(lines))


if __name__ == "__main__":
    main(sys.argv[1])
